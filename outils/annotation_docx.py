"""Pré-annotation du jeu d'évaluation.

  export <docx> <dossier>   : écrit le texte de chaque session (S1.txt, ...) tel qu'il apparaît dans le Word.
  apply <docx> <json> <out> : surligne dans le Word les passages listés dans le JSON.

Format du JSON : [{"id": "S12-T3", "role": "Toi" | "Claude", "color": "jaune" | "vert", "text": "passage exact"}]
"""
import copy
import json
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

LABEL = re.compile(r"^(S\d+)-T(\d+)  ·  (Toi|Claude)")
HEADING = re.compile(r"^(S\d+) — ")
COLORS = {"jaune": WD_COLOR_INDEX.YELLOW, "vert": WD_COLOR_INDEX.BRIGHT_GREEN}


def blocks(doc):
    """Associe chaque paragraphe de contenu au bloc (id du tour, rôle) qui le précède."""
    current = None
    for p in doc.paragraphs:
        m = LABEL.match(p.text)
        if m:
            current = (f"{m.group(1)}-T{m.group(2)}", m.group(3))
            yield ("label", current, p)
            continue
        h = HEADING.match(p.text)
        if h:
            current = None
            yield ("heading", h.group(1), p)
            continue
        if current:
            yield ("content", current, p)


def export(docx_path, out_dir):
    os.makedirs(out_dir, exist_ok=True)
    doc = Document(docx_path)
    files, session = {}, None
    for kind, key, p in blocks(doc):
        if kind == "heading":
            session = key
            files[session] = [p.text]
        elif session and kind == "label":
            files[session].append(f"\n[{key[0]} {key[1]}]")
        elif session and p.text.strip():
            files[session].append(p.text)
    for s, lines in files.items():
        with open(os.path.join(out_dir, f"{s}.txt"), "w", encoding="utf-8") as fh:
            fh.write("\n".join(lines))
    sizes = sorted(((len("\n".join(v)), k) for k, v in files.items()), reverse=True)
    print(f"{len(files)} sessions, {sum(n for n, _ in sizes)} caractères ; plus grosses : {sizes[:5]}")


def highlight(paragraph, start, end, color):
    """Surligne les caractères [start, end) du paragraphe en découpant les runs si besoin."""
    pos = 0
    for run in list(paragraph.runs):
        text = run.text
        r0, r1 = pos, pos + len(text)
        pos = r1
        a, b = max(start, r0), min(end, r1)
        if a >= b:
            continue
        before, middle, after = text[: a - r0], text[a - r0: b - r0], text[b - r0:]
        run.text = middle
        run.font.highlight_color = color
        if before:
            new = copy.deepcopy(run._r)
            run._r.addprevious(new)
            new_run = type(run)(new, paragraph)
            new_run.text = before
            new_run.font.highlight_color = None
        if after:
            new = copy.deepcopy(run._r)
            run._r.addnext(new)
            new_run = type(run)(new, paragraph)
            new_run.text = after
            new_run.font.highlight_color = None


def apply(docx_path, json_path, out_path):
    doc = Document(docx_path)
    with open(json_path, encoding="utf-8") as fh:
        items = json.load(fh)
    by_block = {}
    for kind, key, p in blocks(doc):
        if kind == "content":
            by_block.setdefault(key, []).append(p)
    done, missing = 0, []
    for it in items:
        paragraphs = by_block.get((it["id"], it["role"]), [])
        for p in paragraphs:
            i = p.text.find(it["text"])
            if i >= 0:
                highlight(p, i, i + len(it["text"]), COLORS[it["color"]])
                done += 1
                break
        else:
            missing.append(it)
    doc.save(out_path)
    print(f"{done} surlignages appliqués, {len(missing)} introuvables")
    for it in missing[:30]:
        print("  introuvable :", it["id"], it["role"], repr(it["text"][:60]))


COLOR_NAMES = {v: k for k, v in COLORS.items()}


def read(docx_path, out_json):
    """Relit les surlignages du Word (ceux de l'utilisateur compris), fusionnés par passage continu."""
    doc = Document(docx_path)
    found = []
    for kind, key, p in blocks(doc):
        if kind != "content":
            continue
        span, color = "", None
        for run in p.runs + [None]:
            c = run.font.highlight_color if run is not None else None
            if c != color and span.strip() and color is not None:
                found.append({"id": key[0], "role": key[1], "color": COLOR_NAMES.get(color, str(color)),
                              "text": span.strip()})
            if c != color:
                span, color = "", c
            if run is not None:
                span += run.text
    with open(out_json, "w", encoding="utf-8") as fh:
        json.dump(found, fh, ensure_ascii=False, indent=0)
    print(f"{len(found)} surlignages lus")


if __name__ == "__main__":
    cmd = sys.argv[1]
    if cmd == "read":
        read(sys.argv[2], sys.argv[3])
    if cmd == "export":
        export(sys.argv[2], sys.argv[3])
    elif cmd == "apply":
        apply(sys.argv[2], sys.argv[3], sys.argv[4])
