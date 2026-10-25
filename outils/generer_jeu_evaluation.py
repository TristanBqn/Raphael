"""Génère le document Word du jeu d'évaluation à partir de tout l'historique.

Chaque tour porte un identifiant (S12-T3) pour relier plus tard les surlignages aux échanges.
Usage : python generer_jeu_evaluation.py <fichier.docx>
"""
import re
import sys
from datetime import datetime

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

from historique import load_sessions

MAX_CHARS = 6000   # au-delà, on garde le début (souvent un contenu collé)
KEEP_CHARS = 3000
RED = RGBColor(0xC0, 0x00, 0x00)    # tes messages
BLUE = RGBColor(0x1F, 0x4E, 0x9E)   # les réponses de Claude
GREY = RGBColor(0x6B, 0x72, 0x80)
CTRL = re.compile("[" + "".join(map(chr, [*range(0, 9), 11, 12, *range(14, 32)])) + "]")


def fmt_date(ts: str) -> str:
    try:
        return datetime.fromisoformat(ts.replace("Z", "+00:00")).astimezone().strftime("%d/%m/%Y %H:%M")
    except ValueError:
        return ts


def clip(text: str) -> tuple[str, int]:
    if len(text) <= MAX_CHARS:
        return text, 0
    return text[:KEEP_CHARS], len(text) - KEEP_CHARS


BOLD = re.compile(r"\*\*(.+?)\*\*")


def add_text(doc, text: str, color: RGBColor):
    for line in text.split("\n"):
        line = CTRL.sub("", line)
        heading = line.lstrip().startswith("#")
        if heading:
            line = line.lstrip().lstrip("#").strip()
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        # Les morceaux impairs du découpage sont ceux entourés de ** en Markdown.
        for i, chunk in enumerate(BOLD.split(line)):
            if chunk:
                run = p.add_run(chunk)
                run.bold = heading or i % 2 == 1
                run.font.color.rgb = color


def label(doc, txt: str, color: RGBColor):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(txt)
    r.bold = True
    r.font.color.rgb = color
    r.font.size = Pt(9.5)


def legend_line(doc, color, name: str, meaning: str, example: str):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(name)
    r.bold = True
    r.font.highlight_color = color
    p.add_run(f" : {meaning} Exemple : ")
    r = p.add_run(example)
    r.font.highlight_color = color


def build(out_path: str):
    sessions = load_sessions()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2))
    normal = doc.styles["Normal"]
    normal.font.name = "Segoe UI"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Segoe UI")
    normal.font.size = Pt(10)

    doc.add_heading("Jeu d'évaluation du second cerveau", 0)
    n_user = sum(1 for s in sessions for e in s.events if e.role == "user")
    doc.add_paragraph(f"Tout ton historique Claude Code : {len(sessions)} conversations, {n_user} messages de ta part. "
                      f"Généré le {datetime.now().strftime('%d/%m/%Y')}.")

    doc.add_heading("Comment annoter", 1)
    p = doc.add_paragraph("Tes messages sont en ")
    p.add_run("rouge").font.color.rgb = RED
    p.add_run(", les réponses de Claude en ")
    p.add_run("bleu").font.color.rgb = BLUE
    p.add_run(".")
    doc.add_paragraph("Lis les échanges et surligne directement dans le texte, avec l'outil Surbrillance de Word "
                      "(onglet Accueil). Deux couleurs seulement :")
    legend_line(doc, WD_COLOR_INDEX.YELLOW, "Jaune — à retenir",
                "les mots ou la phrase qui auraient dû faire enregistrer un souvenir.",
                "Mon NAS est un Synology installé au bureau")
    legend_line(doc, WD_COLOR_INDEX.BRIGHT_GREEN, "Vert — à aller chercher",
                "les mots qui auraient dû faire consulter le second cerveau avant de répondre.",
                "sauvegarde ce dossier sur mon NAS")
    tips = [
        "Surligne le moins possible : seulement les mots qui déclenchent, pas tout le paragraphe.",
        "Tu peux surligner aussi dans les réponses de Claude et dans les lignes d'erreur d'outil.",
        "Facultatif : ajoute un commentaire Word pour préciser le type (fait, préférence, procédure, "
        "leçon d'échec, historique de projet) ou ce qu'il aurait fallu retrouver.",
        "Ne modifie pas les identifiants comme S12-T3 : ils servent à relier tes surlignages aux échanges.",
        "Tu n'es pas obligé de tout annoter. Mieux vaut 50 échanges bien annotés que 500 survolés.",
    ]
    for t in tips:
        doc.add_paragraph(t, style="List Bullet")
    doc.add_paragraph("Les secrets (clés API, mots de passe, jetons) ont été remplacés par [SECRET MASQUÉ]. "
                      f"Les messages de plus de {MAX_CHARS} caractères, souvent des contenus collés, sont coupés "
                      "après leurs premiers caractères, et la coupure est signalée.")

    for si, s in enumerate(sessions, 1):
        doc.add_page_break()
        doc.add_heading(f"S{si} — {fmt_date(s.started)} — {s.project}", 1)
        turn = 0
        last_label = None
        for e in s.events:
            if e.role == "user":
                turn += 1
                label(doc, f"S{si}-T{turn}  ·  Toi  ·  {fmt_date(e.timestamp)}", RED)
                last_label = "user"
            elif e.role == "assistant":
                # Une seule étiquette pour les morceaux de réponse consécutifs d'un même tour.
                if last_label != "assistant":
                    label(doc, f"S{si}-T{turn}  ·  Claude", BLUE)
                last_label = "assistant"
            else:
                p = doc.add_paragraph()
                r = p.add_run(f"⚠ Erreur d'outil — {e.text}")
                r.italic = True
                r.font.size = Pt(8.5)
                r.font.color.rgb = GREY
                continue
            text, cut = clip(e.text)
            add_text(doc, text, RED if e.role == "user" else BLUE)
            if cut:
                p = doc.add_paragraph()
                r = p.add_run(f"[… {cut} caractères coupés …]")
                r.italic = True
                r.font.color.rgb = GREY

    doc.save(out_path)
    print(f"{len(sessions)} sessions, {n_user} messages → {out_path}")


if __name__ == "__main__":
    build(sys.argv[1])
